from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

doc = Document()

# Title
title = doc.add_heading('AMAS — Autonomous Multi-Agent Task Automation System', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Subtitle
subtitle = doc.add_heading('Internship Project Submission Report', level=1)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Author info
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Author: Aman Kumar')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(100, 116, 139)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Technology Stack: Python 3.10+, TypeScript/React 19, SQLite, Vite, Tailwind CSS')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(100, 116, 139)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Live Deployment: https://amas-autonomus-multi-agent-task-automation-syste-production.up.railway.app/')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(100, 116, 139)
run.font.underline = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('License: MIT')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(100, 116, 139)

doc.add_page_break()

# Table of Contents
doc.add_heading('Table of Contents', level=1)
toc_items = [
    '1. Executive Summary',
    '2. Problem Statement',
    '3. Solution Architecture',
    '4. Core Components',
    '5. Frontend: Workflow Studio',
    '6. CLI & REST API Reference',
    '7. Quality Benchmarks & Evaluation Suite',
    '8. Repository Structure',
    '9. Key Technical Achievements',
    '10. Quick Start Guide',
    '11. Open Source Acknowledgments',
    '12. License',
    '13. Conclusion'
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.style = 'List Number'

doc.add_page_break()

# 1. Executive Summary
doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph(
    'AMAS (Autonomous Multi-Agent Task Automation System) is an enterprise-grade agentic workflow orchestration platform '
    'that transforms standard prompt-and-response AI into an autonomous execution engine. The system dynamically plans, '
    'executes, audits, and self-heals complex analytical and operational tasks through a sophisticated multi-agent architecture.'
)

doc.add_heading('Key Innovation', level=2)
doc.add_paragraph(
    'Unlike fragile LLM wrappers or monolithic chat interfaces that hallucinate calculations and fabricate data, AMAS implements:'
)
innovations = [
    'Dynamic Directed Acyclic Graph (DAG) task synthesis',
    'Strict 5-tier tool hierarchy precedence',
    'Shared typed blackboard memory',
    'Independent verification gates with mathematical recalculation',
    'Multi-provider web intelligence',
    'Provider-agnostic model routing'
]
for item in innovations:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# 2. Problem Statement
doc.add_heading('2. Problem Statement', level=1)
doc.add_paragraph(
    'Modern LLM applications frequently fail in production due to three architectural flaws:'
)

problems = [
    ('Monolithic Hallucination', 
     'Asking a single conversational model to do research, solve math, format code, and verify its own output leads to fabrication of facts and erroneous calculations.'),
    ('Brittle Rigid Routing', 
     'Hardcoding rule engines breaks as soon as a user asks multi-domain questions.'),
    ('Context Degradation & Isolation', 
     'Multi-agent pipelines often pass unstructured conversational snippets between agents like a game of telephone, losing precision, table data, and citations.')
]

for title_text, desc in problems:
    p = doc.add_paragraph()
    run = p.add_run(f'{title_text}: ')
    run.bold = True
    p.add_run(desc)

doc.add_page_break()

# 3. Solution Architecture
doc.add_heading('3. Solution Architecture', level=1)

doc.add_heading('3.1 System Architecture Layers', level=2)
doc.add_paragraph(
    'The AMAS architecture follows a layered design with clear separation of concerns:'
)

layers = [
    ('Presentation Layer', 'React 19 + TypeScript + Vite + Tailwind CSS | Workflow Studio | Real-time SSE Stream'),
    ('API Gateway & Express Server', 'REST Endpoints | Process Isolation | UTF-8 Streaming | JSON Boundary Parser'),
    ('AMAS Control Plane', 'RunManager | TaskGraph (DAG DFS) | PolicyManager | EventBus | RecoveryAgent | VerificationAuditor'),
    ('AMAS Runtime & Agent Engine', 'RunMemory (Scoped Blackboard) | Dynamic DAG Planner | PraisonRuntime | Argument Sanitizer'),
    ('Provider-Agnostic LLM Layer', 'Groq | OpenRouter | CodeCraft | Gemini | Local (Ollama/vLLM) | Health-Aware Routing'),
    ('Multi-Tier Tool Ecosystem', '5-Tier Precedence: PraisonAI -> LangChain -> MCP -> Official SDKs -> Custom Tools'),
    ('Persistence & Storage', 'SQLite3 (amas.db) | Artifact Files | Tool Metrics')
]

for layer_name, description in layers:
    p = doc.add_paragraph()
    run = p.add_run(f'{layer_name}: ')
    run.bold = True
    p.add_run(description)

doc.add_page_break()

# 4. Core Components
doc.add_heading('4. Core Components', level=1)

# 4.1 Control Plane
doc.add_heading('4.1 Control Plane & DAG Orchestration Engine', level=2)

doc.add_heading('Dynamic DAG Planning (amas/runtime/planner.py)', level=3)
dag_features = [
    'LLM-driven task graph generation from natural language objectives',
    'Topological dependencies with explicit task IDs',
    'Least-privilege tool assignment per task',
    'Anti-placeholder validation (prevents template strings like {region}, {date})',
    'Cycle detection via DFS validation'
]
for item in dag_features:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Shared Blackboard Memory (amas/runtime/memory.py)', level=3)
memory_features = [
    'Thread-safe typed blackboard for inter-agent communication',
    'completed_task_results: Automatic recording of all task outputs',
    'get_accumulated_results_text(): Concatenates upstream findings',
    'Key-value store for numerical series, tickers, timestamps',
    'Deliverable aggregation via RunManager._build_latest_computation()'
]
for item in memory_features:
    doc.add_paragraph(item, style='List Bullet')

# 4.2 Runtime Engine
doc.add_heading('4.2 Runtime Engine & PraisonAI Integration', level=2)
doc.add_paragraph(
    'The Runtime Engine (amas/runtime/praison_runtime.py) connects the control plane to agent execution:'
)
runtime_features = [
    'Leverages praisonaiagents for agent role configuration, goal framing, and tool attachment',
    'Seamless fallback to internal execution when PraisonAI unavailable',
    'Intelligent argument resolution and sanitization',
    'Search query normalization (cleans templates, removes restrictive filters)',
    'Artifact writer automation (injects accumulated research)',
    'Adaptive context windowing (up to 4,000 characters)'
]
for item in runtime_features:
    doc.add_paragraph(item, style='List Bullet')

# 4.3 Six Specialized Agent Roles
doc.add_heading('4.3 Six Specialized Agent Roles', level=2)
doc.add_paragraph(
    'AMAS enforces strict role separation across 6 specialized agents (amas/runtime/agents.py):'
)

# Create table for agents
table = doc.add_table(rows=7, cols=4)
table.style = 'Light Grid Accent 1'

# Header row
headers = ['Agent Role', 'System Prompt Directive', 'Permitted Tools', 'Core Responsibility']
for i, header in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = header
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt(9)

# Data rows
agents_data = [
    ['PlanningAgent', 'Architect and decompose objectives into minimum sequential tasks.', 'None (pure reasoning)', 'Formulates DAG, sets dependencies, enforces least-privilege scoping'],
    ['ResearchAgent', 'Gather authentic facts from live search, encyclopedias, and market feeds.', 'Web search, Tavily, Financial retriever, Wikipedia', 'Live data retrieval without fabricating facts'],
    ['AnalysisAgent', 'Perform deterministic mathematical and statistical operations.', 'Risk calculator, Anomaly detector, Code interpreter', 'Computes Sharpe ratios, volatility, Z-scores'],
    ['ExecutionAgent', 'Synthesize deliverable reports and export structured documents.', 'Artifact writer, File read', 'Compiles comprehensive markdown deliverables'],
    ['VerificationAgent', 'Adversarially audit claims, re-run math, and verify evidence coverage.', 'Independent calculation methods', 'Validates bounds, verifies citations, scores quality'],
    ['RecoveryAgent', 'Diagnose failures, select healing strategies, and restore execution.', 'Control plane policy operations', 'Self-healing via provider switching, replanning']
]

for row_idx, agent_data in enumerate(agents_data, 1):
    for col_idx, cell_text in enumerate(agent_data):
        cell = table.rows[row_idx].cells[col_idx]
        cell.text = cell_text
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(8)

doc.add_page_break()

# 4.4 Provider-Agnostic LLM Routing
doc.add_heading('4.4 Provider-Agnostic LLM Routing', level=2)

doc.add_heading('Supported Providers', level=3)
providers = [
    ('Groq', 'Ultra-high throughput inference (openai/gpt-oss-120b, qwen/qwen3.8-27b)'),
    ('OpenRouter', 'Access to 100+ frontier models (llama-3.3-70b, claude-3.5-sonnet, deepseek-chat)'),
    ('CodeCraft', 'High-throughput code generation and structured JSON output'),
    ('Google Gemini', 'Deep context reasoning (gemini-2.5-flash, gemini-2.5-pro)'),
    ('Local OpenAI Compatible', 'Local inference via Ollama or vLLM (http://localhost:11434/v1)')
]
for name, desc in providers:
    p = doc.add_paragraph()
    run = p.add_run(f'{name}: ')
    run.bold = True
    p.add_run(desc)

doc.add_heading('Routing Modes', level=3)
modes = [
    ('Fixed', 'Strictly routes all calls to the user-selected provider'),
    ('Fallback', 'Tries primary provider; automatically fails over to secondary upon exhaustion or HTTP 429'),
    ('Automatic (Health-Aware)', 'Dynamically selects healthiest provider based on success rate, latency EMA, rate-limit cooldown')
]
for name, desc in modes:
    p = doc.add_paragraph()
    run = p.add_run(f'{name}: ')
    run.bold = True
    p.add_run(desc)

# 4.5 Multi-Tier Tool Ecosystem
doc.add_heading('4.5 Multi-Tier Tool Ecosystem (5-Tier Precedence)', level=2)
doc.add_paragraph(
    'To avoid code sprawl and enforce software reuse, AMAS organizes all tools into a strict 5-Tier Precedence Hierarchy:'
)

tiers = [
    ('Tier 1', 'PraisonAI Official Ecosystem (Web Search, File Read, Code Interpreter)'),
    ('Tier 2', 'LangChain Community Tools (Wikipedia Retriever)'),
    ('Tier 3', 'Model Context Protocol (MCP) Tools'),
    ('Tier 4', 'Official Provider & Service SDKs (Tavily, Yahoo Finance, Exa, Brave, Serper)'),
    ('Tier 5', 'AMAS Custom Tools (Risk Calculator, Anomaly Detector, Artifact Writer)')
]
for tier, desc in tiers:
    p = doc.add_paragraph()
    run = p.add_run(f'{tier}: ')
    run.bold = True
    p.add_run(desc)

doc.add_page_break()

# 4.6 Independent Verification
doc.add_heading('4.6 Independent Verification & Adversarial Audit Gate', level=2)
doc.add_paragraph(
    'To eliminate LLM self-grading bias, the Verification Auditor (amas/verification/auditor.py) runs objective assertions:'
)

verification_items = [
    ('Independent Numerical Recalculation', 
     'Financial: Recalculates daily returns, annualized volatility, Sharpe ratio from raw prices (5% tolerance). '
     'Anomaly: Re-computes standard deviations and Z-scores against raw sensor vectors.'),
    ('Mathematical Boundary Checks', 
     'Sharpe ratio bounds: [-10, 25]; Non-negative volatility; Probability values <= 1.0'),
    ('Evidence & Citation Coverage', 
     'Ratio of factual assertions supported by search citations; URL validity and source diversity'),
    ('Weighted Quality Score', 
     'Score = 0.40 x NumericalAccuracy + 0.30 x EvidenceCoverage + 0.20 x CitationCoverage + 0.10 x Completeness. '
     '>= 85% receives VERIFIED certification badge.')
]

for title_text, desc in verification_items:
    p = doc.add_paragraph()
    run = p.add_run(f'{title_text}: ')
    run.bold = True
    p.add_run(desc)

# 4.7 Self-Healing
doc.add_heading('4.7 Self-Healing & Diagnostic Recovery', level=2)
doc.add_paragraph(
    'When a task fails or verification rejects an output, the Recovery Agent (amas/control_plane/recovery_agent.py) '
    'diagnoses the root cause and selects an appropriate recovery strategy:'
)

recovery_strategies = [
    ('Rate Limit / 429', 'SWITCH_PROVIDER'),
    ('Search Network Error', 'SWITCH_SEARCH_PROVIDER'),
    ('Invalid Tool Arguments', 'MODIFY_ARGUMENTS'),
    ('Math Verification Fail', 'REPLAN_TASK'),
    ('Timeout / Glitch', 'RETRY_WITH_BACKOFF'),
    ('Sensitive Operation', 'REQUEST_HUMAN_APPROVAL'),
    ('Permission / Violation', 'ABORT')
]

table2 = doc.add_table(rows=len(recovery_strategies)+1, cols=2)
table2.style = 'Light Grid Accent 1'
table2.rows[0].cells[0].text = 'Fault Type'
table2.rows[0].cells[1].text = 'Recovery Strategy'
for cell in table2.rows[0].cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, (fault, strategy) in enumerate(recovery_strategies, 1):
    table2.rows[i].cells[0].text = fault
    table2.rows[i].cells[1].text = strategy

doc.add_paragraph()
recovery_features = [
    'Exponential backoff with jitter (+/-25%) prevents hammering endpoints',
    'Max retries bound (default: 3) prevents infinite loops',
    'Full audit logging to database with error classifications and rationale'
]
for item in recovery_features:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# 4.8 Security Architecture
doc.add_heading('4.8 Security Architecture & Restricted Execution', level=2)
security_items = [
    'Arbitrary RCE Blocked: /api/run-python returns HTTP 403 for arbitrary scripts',
    'Restricted Python Evaluator: Dangerous built-ins stripped (__import__, eval, exec, open, compile); Critical modules blocked (os, sys, subprocess, shutil, socket); Only safe math operations allowed',
    'Workspace Path Isolation: Directory traversal blocked; file operations restricted to project subdirectories',
    'Tool Permission Boundaries: Granular permissions (read_only, workspace_bound, pure_math, requires_approval)',
    'Human-in-the-Loop Approvals: High-impact operations require explicit approval via SQLite states (PENDING -> APPROVED/REJECTED)',
    'Zero Secret Exposure: API keys never returned in REST endpoints, written to SQLite, logged, or sent over SSE'
]
for item in security_items:
    doc.add_paragraph(item, style='List Bullet')

# 4.9 Persistence
doc.add_heading('4.9 Persistence, SQLite Schema & Storage Layer', level=2)
doc.add_paragraph('All workflow operations persist to SQLite (amas/storage/amas.db) via amas/storage/database.py:')

tables = [
    ('runs', 'Workflow executions with metrics, blackboard state, provider/model info'),
    ('tasks', 'Individual task nodes with results, verification, duration, dependencies'),
    ('events', 'Real-time event log for SSE streaming (RUN_CREATED, TASK_STARTED, VERIFICATION_COMPLETED, etc.)'),
    ('tool_calls', 'Execution history with latency metrics, success/failure counts'),
    ('artifacts', 'Generated deliverable files (markdown reports, analytical summaries)'),
    ('approvals', 'Human approval workflow states with durable SQLite persistence'),
    ('tool_metrics', 'Persistent performance telemetry: calls, successes, failures, avg latency per tool')
]

table3 = doc.add_table(rows=len(tables)+1, cols=2)
table3.style = 'Light Grid Accent 1'
table3.rows[0].cells[0].text = 'Table'
table3.rows[0].cells[1].text = 'Description'
for cell in table3.rows[0].cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, (name, desc) in enumerate(tables, 1):
    table3.rows[i].cells[0].text = name
    table3.rows[i].cells[1].text = desc

doc.add_page_break()

# 5. Frontend
doc.add_heading('5. Frontend: Workflow Studio', level=1)
doc.add_paragraph(
    'The frontend is built on React 19, TypeScript, and Tailwind CSS, designed for maximum clarity, '
    'real-time observability, and actionable intelligence.'
)

doc.add_heading('Interactive Scenario Presets', level=2)
presets = [
    ('India Floods & Relief', 'Real-time casualty tracking, SDRF metrics, infrastructure damage breakdown'),
    ('IPL & Global Sports', 'Live sports wire retrieval of match outcomes, team scorelines, player MVPs'),
    ('Oscars & Box Office', 'Global revenues, Academy Awards won, Rotten Tomatoes/Metacritic consensus'),
    ('SpaceX Starship Telemetry', 'Flight telemetry, stage separation status, Raptor engine performance'),
    ('Stock Risk & Math', 'Live 30-day volatility and Sharpe ratio computation for equities (AAPL, NVDA)'),
    ('Telemetry Anomaly Detection', 'Statistical outlier and sensor spike detection with self-healing demo')
]
for name, desc in presets:
    p = doc.add_paragraph()
    run = p.add_run(f'{name}: ')
    run.bold = True
    p.add_run(desc)

doc.add_heading('Key Features', level=2)
frontend_features = [
    'Live DAG Execution Visualizer: Displays task nodes, dependencies, status (PENDING/RUNNING/COMPLETED/FAILED/VERIFIED), agent roles, duration',
    'Custom GFM Markdown Renderer: Full GitHub Flavored Markdown parsing with formatted tables, syntax highlighting, citations',
    'Domain-Specific Deliverable Displays: Disaster (stat cards, regional breakdown), Financial (risk metrics, verified math), Anomaly (Z-scores, outliers), Aerospace (flight milestones, telemetry), Sports/Entertainment',
    'Real-Time SSE Event Stream: Live agent logs, verification audit events, self-healing transitions',
    'Provider Settings Modal: Switch active LLM providers, edit endpoints, real-time latency ping tests',
    'PDF/DOCX Export: Professional report generation with audit certificates and formatted metrics'
]
for item in frontend_features:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# 6. CLI & REST API
doc.add_heading('6. CLI & REST API Reference', level=1)

doc.add_heading('CLI Commands', level=2)
cli_commands = [
    ('python -m amas.cli run "Analyze AAPL 30-day Sharpe ratio"', 'Run an autonomous workflow'),
    ('python -m amas.cli providers', 'List registered providers and health metrics'),
    ('python -m amas.cli test-provider groq', 'Test connectivity and latency for a provider'),
    ('python -m amas.cli tools', 'List all registered tools and source tiers'),
    ('python -m amas.cli execute-tool official_financial_retriever \'{"ticker": "AAPL", "period_days": 14}\'', 'Execute a tool directly with JSON input'),
    ('python -m amas.cli eval', 'Run the automated evaluation suite')
]

table4 = doc.add_table(rows=len(cli_commands)+1, cols=2)
table4.style = 'Light Grid Accent 1'
table4.rows[0].cells[0].text = 'Command'
table4.rows[0].cells[1].text = 'Description'
for cell in table4.rows[0].cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt(9)

for i, (cmd, desc) in enumerate(cli_commands, 1):
    table4.rows[i].cells[0].text = cmd
    table4.rows[i].cells[1].text = desc
    for paragraph in table4.rows[i].cells[0].paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(8)
            run.font.name = 'Consolas'

doc.add_heading('Key REST Endpoints', level=2)
endpoints = [
    ('/api/system-info', 'GET', 'System health, active provider, model, capabilities'),
    ('/api/providers', 'GET', 'Provider catalog with health stats'),
    ('/api/providers/:id/test', 'POST', 'Live connection test for a provider'),
    ('/api/tools', 'GET', 'Tool registry with sources, permissions, metrics'),
    ('/api/tools/execute', 'POST', 'Execute a tool directly under permission boundaries'),
    ('/api/workflow/run', 'POST', 'Execute autonomous multi-agent DAG workflow'),
    ('/api/runs/:runId/events', 'GET', 'Real-time SSE stream for live DAG updates'),
    ('/api/evaluation/run', 'POST', 'Run automated benchmark test suite')
]

table5 = doc.add_table(rows=len(endpoints)+1, cols=3)
table5.style = 'Light Grid Accent 1'
for i, header in enumerate(['Endpoint', 'Method', 'Description']):
    table5.rows[0].cells[i].text = header
    for paragraph in table5.rows[0].cells[i].paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, (endpoint, method, desc) in enumerate(endpoints, 1):
    table5.rows[i].cells[0].text = endpoint
    table5.rows[i].cells[1].text = method
    table5.rows[i].cells[2].text = desc

doc.add_page_break()

# 7. Quality Benchmarks
doc.add_heading('7. Quality Benchmarks & Evaluation Suite', level=1)
doc.add_paragraph(
    'AMAS features an automated evaluation benchmark (python -m amas.cli eval) testing 7 operational dimensions:'
)

benchmarks = [
    'Dynamic DAG Planning: Asserts task graphs are dynamically synthesized without hardcoded routing',
    'Cycle Prevention: Validates topological DFS sorting and cyclic graph rejection',
    'Tool Precedence Compliance: Confirms strict adherence to the 5-Tier Tool Hierarchy',
    'Provider Failover: Verifies automatic failover from rate-limited primary providers',
    'Mathematical Verification: Asserts corrupted calculation outputs (e.g. Sharpe ratio 150) are rejected',
    'Search Multi-Provider Failover: Validates fallback from primary search SDKs to alternatives',
    'Sandbox Security Boundaries: Confirms unauthorized builtins and system calls are blocked'
]
for item in benchmarks:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph(
    'Test Results: All 7 operational dimensions validated with 100% pass rate when API keys configured.'
)

# 8. Repository Structure
doc.add_heading('8. Repository Structure', level=1)
doc.add_paragraph('The project follows a clean modular architecture:')

structure = [
    ('amas/', 'Core AMAS Engine Package'),
    ('  control_plane/', 'DAG orchestration & execution engine (event_bus, policy_manager, run_manager, task_graph, recovery_agent)'),
    ('  providers/', 'Provider-Agnostic LLM Layer (base, gemini_provider, manager, openai_compatible)'),
    ('  runtime/', 'Agent Runtime Engine (agents, memory, planner, praison_runtime)'),
    ('  storage/', 'Persistent Storage & Database (artifacts, amas.db, database.py)'),
    ('  tools/', 'Multi-Tier Tool Ecosystem (adapters, custom, official, web_search_manager, base, registry)'),
    ('  verification/', 'Independent Verification & Audit (auditor.py)'),
    ('  cli.py', 'Unified AMAS Command-Line Interface'),
    ('server.ts', 'Express API server with SSE endpoint & Vite dev middleware'),
    ('src/', 'React 19 + TypeScript Frontend (components, utils, App.tsx, types.ts)'),
    ('autonomous_multi_agent_system.py', 'Standalone CLI entrypoint (legacy compatibility)'),
    ('tests/', 'Automated Unit & Integration Tests (test_amas_suite.py)'),
    ('requirements.txt', 'Python dependencies'),
    ('package.json', 'Frontend dependencies & scripts'),
    ('.env.example', 'Environment variable configuration template'),
    ('Dockerfile', 'Container configuration'),
    ('README.md', 'Comprehensive documentation')
]

for path, desc in structure:
    p = doc.add_paragraph()
    run = p.add_run(f'{path}')
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.bold = True
    p.add_run(f' - {desc}')

doc.add_page_break()

# 9. Key Technical Achievements
doc.add_heading('9. Key Technical Achievements', level=1)

achievements = [
    ('Zero Hardcoded Routing', 'LLM-driven dynamic DAG compilation replaces keyword-based routing'),
    ('Strict Tool Hierarchy', '5-tier precedence with source attribution prevents tool sprawl'),
    ('Mathematical Verification', 'Independent recalculation with 5% tolerance eliminates hallucination'),
    ('Multi-Provider Resilience', 'Health-aware routing with automatic fallback ensures uptime'),
    ('Self-Healing Execution', 'Classified recovery strategies beyond simple retry'),
    ('Security by Design', 'Sandboxed execution, path isolation, approval gates, zero secret exposure'),
    ('Production Observability', 'SSE streaming, SQLite persistence, tool metrics, audit trails'),
    ('Professional Deliverables', 'PDF/DOCX export with audit certificates and formatted metrics')
]

table6 = doc.add_table(rows=len(achievements)+1, cols=2)
table6.style = 'Light Grid Accent 1'
table6.rows[0].cells[0].text = 'Achievement'
table6.rows[0].cells[1].text = 'Implementation'
for cell in table6.rows[0].cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, (ach, impl) in enumerate(achievements, 1):
    table6.rows[i].cells[0].text = ach
    table6.rows[i].cells[1].text = impl

doc.add_page_break()

# 10. Quick Start Guide
doc.add_heading('10. Quick Start Guide', level=1)

doc.add_heading('Prerequisites', level=2)
prereqs = ['Node.js v18.0.0 or higher', 'Python 3.10, 3.11, or 3.12']
for item in prereqs:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Installation', level=2)
doc.add_paragraph('Clone and install dependencies:')
code = doc.add_paragraph()
run = code.add_run('git clone https://github.com/AmanKumar-St/AMAS-Autonomus-Multi-agent-Task-Automation-system.git\ncd AMAS-Autonomus-Multi-agent-Task-Automation-system\n\n# Install Node frontend dependencies\nnpm install\n\n# Install Python backend dependencies\npip install -r requirements.txt')
run.font.name = 'Consolas'
run.font.size = Pt(9)

doc.add_paragraph('Configure environment variables:')
code2 = doc.add_paragraph()
run2 = code2.add_run('cp .env.example .env\n# Edit .env with your API keys:\n# GROQ_API_KEY="gsk_..."\n# TAVILY_API_KEY="tvly_..."\n# OPENROUTER_API_KEY="sk-or-..."\n# GEMINI_API_KEY="..."')
run2.font.name = 'Consolas'
run2.font.size = Pt(9)

doc.add_paragraph('Launch the application:')
code3 = doc.add_paragraph()
run3 = code3.add_run('npm run dev')
run3.font.name = 'Consolas'
run3.font.size = Pt(9)
doc.add_paragraph('Navigate to http://localhost:3000 in your browser.')

doc.add_heading('Run Automated Test Suite', level=2)
code4 = doc.add_paragraph()
run4 = code4.add_run('python -m unittest tests/test_amas_suite.py -v')
run4.font.name = 'Consolas'
run4.font.size = Pt(9)

doc.add_page_break()

# 11. Open Source Acknowledgments
doc.add_heading('11. Open Source Acknowledgments', level=1)
doc.add_paragraph(
    'AMAS uses and integrates concepts and tooling from the open-source PraisonAI project. '
    'This repository acknowledges PraisonAI for its agent orchestration patterns, tool integration ideas, '
    'and runtime approach that informed parts of the AMAS execution stack.'
)

acknowledgments = [
    'PraisonAI: https://praison.ai',
    'PraisonAI GitHub: https://github.com/praisonai'
]
for item in acknowledgments:
    doc.add_paragraph(item, style='List Bullet')

# 12. License
doc.add_heading('12. License', level=1)
doc.add_paragraph(
    'This project is licensed under the MIT License. See the LICENSE file for details.'
)

# 13. Conclusion
doc.add_heading('13. Conclusion', level=1)
doc.add_paragraph(
    'AMAS represents a significant advancement in autonomous multi-agent systems by solving the fundamental '
    'problems of LLM production deployments: hallucination, brittle routing, and context degradation. Through its '
    'novel architecture combining dynamic DAG planning, strict tool governance, independent verification, and '
    'self-healing recovery, AMAS delivers production-ready autonomous task automation with mathematical rigor '
    'and operational observability.'
)

doc.add_paragraph(
    'The system demonstrates that AI agents can be engineered with the same reliability principles as traditional '
    'software systems through separation of concerns, independent auditing, structured error handling, and '
    'defense-in-depth security while maintaining the flexibility and intelligence that makes LLM-based automation valuable.'
)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('End of Report')
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(16, 185, 129)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('AMAS — Autonomous Multi-Agent Task Automation System\n(c) 2026 Aman Kumar')
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(148, 163, 184)

# Save
doc.save('AMAS_Project_Report.docx')
print('DOCX report generated successfully!')